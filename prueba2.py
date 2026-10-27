# -*- coding: utf-8 -*-
"""Aplicación web Streamlit para seguimiento de estudiantes desertores."""
import io
import re
import json
import difflib
import datetime
import unicodedata
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

APP_TITLE = "GESTOR DE SEGUIMIENTO DE ESTUDIANTES DESERTORES"
SUBTITULO = "Registro y control de la búsqueda activa y focalizada - SIMAT"
PIE = "Creado por: Jaime Hernandez - Montería"
LEMA = "La educación nos une, el futuro nos inspira"
SUFIJO = "_seguimineto_desertores"
NOMBRE_GENERAL = "consolidado" + SUFIJO + ".xlsx"
TODAS = "Todas las instituciones"
HEADERS = [
    "PER_ID", "ANNO_INF_2025", "CODIGO_DANE_2025", "NOMBRE_ESTABLECIMIENTO_2025",
    "CODIGO_DANE_SEDE_2025", "NOMBRE_SEDE_2025", "TIPO_DOCUMENTO_2025",
    "NRO_DOCUMENTO_2025", "APELLIDO1_2025", "APELLIDO2_2025", "NOMBRE1_2025",
    "NOMBRE2_2025", "CTE_ID_SECTOR_2025", "ESTADO_DEFINITIVO_2025",
    "Situación actual del estudiante en SIMAT",
    "Ha realizado anteriormente búsqueda activa y focalizada con relación a este estudiante",
    "Detalle BREVEMENTE la gestión realizada (con anterioridad a este reporte) de búsqueda activa y focalizada de este estudiante (indique si la tiene documentada)",
    "Realizó búsqueda activa y focalizada producto del presente requerimiento con relación a este estudiante",
    "Detalle BREVEMENTE la gestión realizada actualmente (producto del presente reporte) para la búsqueda activa y focalizada de este estudiante detallado como posible población desescolarizada (indique si la tiene documentada)",
    "En caso de NO realizar búsqueda activa y focalizada consigne aquí la justificación y argumento (brevemente)",
    "Conclusión de la gestión de búsqueda activa y focalizada",
    "Causa identificada de la situación de retiro de la IE",
    "En caso de marcar otra en la columna anterior - aquí detallar cual es la otra causa (de manera breve)",
    "Activación de las rutas de gestión necesarias para favorecer su permanencia o retorno al sistema educativo.",
    "EN CASO DE ACTIVACIÓN: Detalle brevemente las rutas activadas",
    "EN CASO DE NO ACTIVACIÓN: Detalle brevemente la justificación",
]
N_INFO = 14                      # columnas A:N (solo consulta)
COLS_TEXTO = (2, 4)              # códigos DANE: siempre como texto
SI_NO = ["Sí", "No"]

# Campos editables O:Z (índice 14..25)
CAMPOS = [
    dict(i=14, corto="Situación actual en SIMAT", tipo="combo", cerrado=False, alto=0,
         op=["Matriculado", "Retirado", "Desertor", "Graduado", "Trasladado", "Sin información", "Otra"]),
    dict(i=15, corto="Búsqueda activa y focalizada (anterior)", tipo="combo", cerrado=True, alto=0,
         op=SI_NO + ["No se tiene información"]),
    dict(i=16, corto="Detalle (anterior)", tipo="texto", alto=4),
    dict(i=17, corto="Búsqueda activa y focalizada (actual)", tipo="combo", cerrado=True, alto=0, op=SI_NO),
    dict(i=18, corto="Detalle (actual)", tipo="texto", alto=4),
    dict(i=19, corto="Justificación si no realiza búsqueda", tipo="texto", alto=4),
    dict(i=20, corto="Conclusión de la gestión", tipo="combo", cerrado=False, alto=0,
         op=["Estudiante localizado y matriculado", "Estudiante localizado, pendiente de matrícula",
             "Estudiante localizado, no desea retornar", "Continúa desescolarizado",
             "No fue posible localizarlo", "Trasladado a otra IE o municipio", "No aplica", "Otra"]),
    dict(i=21, corto="Causa de retiro de la IE", tipo="combo", cerrado=False, alto=0,
         op=["Cambio de domicilio", "Cambio de institución", "Dificultades económicas",
             "Problemas familiares", "Desinterés del estudiante o acudiente", "Trabajo infantil",
             "Problemas de salud", "Migración", "Embarazo adolescente", "Conflicto o violencia",
             "No se logró establecer contacto", "No identificada", "Otra"]),
    dict(i=22, corto="Otra causa (si aplica)", tipo="texto", alto=4),
    dict(i=23, corto="Activación de rutas de gestión", tipo="combo", cerrado=True, alto=0, op=SI_NO),
    dict(i=24, corto="Detalle rutas activadas", tipo="texto", alto=5),
    dict(i=25, corto="Detalle si no se activa", tipo="texto", alto=5),
]
LETRAS = [get_column_letter(i + 1) for i in range(26)]


def sin_tildes(t):
    t = unicodedata.normalize("NFD", str(t or ""))
    return "".join(c for c in t if unicodedata.category(c) != "Mn")


def norm(t):
    return re.sub(r"\s+", " ", sin_tildes(t)).strip().casefold()


def norm_enc(t):
    return norm(t).strip(" .:;")


def slug(t):
    s = re.sub(r"[^a-z0-9]+", "", sin_tildes(t).lower())
    return s or "institucion"


def a_texto(v):
    """Convierte cualquier valor de celda a texto limpio (DANE y documentos sin '.0')."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d") if v.time() == datetime.time(0) else v.isoformat(sep=" ")
    if isinstance(v, datetime.date):
        return v.isoformat()
    return str(v).strip()


def limpiar(v):
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def estado_seguimiento(ed):
    """0 = pendiente (sin datos), 1 = en proceso, 2 = diligenciado completo."""
    g = lambda i: str(ed.get(i, "") or "").strip()
    if not any(g(i) for i in range(14, 26)):
        return 0
    obligatorios = [14, 15, 17, 20, 21, 23]
    ok = all(g(i) for i in obligatorios)
    si = lambda i: norm(g(i)) in ("si", "s")
    no = lambda i: norm(g(i)) == "no"
    if si(15):
        ok = ok and bool(g(16))
    if si(17):
        ok = ok and bool(g(18))
    if no(17):
        ok = ok and bool(g(19))
    if norm(g(21)).startswith("otra"):
        ok = ok and bool(g(22))
    if si(23):
        ok = ok and bool(g(24))
    if no(23):
        ok = ok and bool(g(25))
    return 2 if ok else 1


def buscar_fila_encabezado(filas):
    esperados = {norm_enc(h) for h in HEADERS[:N_INFO]}
    for idx in range(min(15, len(filas))):
        celdas = {norm_enc(x) for x in filas[idx] if x is not None}
        if len(esperados & celdas) >= 4:
            return idx
    return None


def mapear_columnas(enc):
    """Asocia cada encabezado esperado con la columna real del archivo (por nombre)."""
    n = len(enc)
    enc_n = [norm_enc(x) for x in enc]
    mapa, usados, avisos = {}, set(), []
    for i, h in enumerate(HEADERS):                       # 1) coincidencia exacta
        hn = norm_enc(h)
        for j, e in enumerate(enc_n):
            if j not in usados and e == hn:
                mapa[i] = j
                usados.add(j)
                break
    for i, h in enumerate(HEADERS):                       # 2) coincidencia aproximada
        if i in mapa:
            continue
        hn = norm_enc(h)
        mejor, mj = 0.0, None
        for j, e in enumerate(enc_n):
            if j in usados or not e:
                continue
            r = difflib.SequenceMatcher(None, hn, e).ratio()
            if r > mejor:
                mejor, mj = r, j
        if mj is not None and mejor >= 0.85:
            mapa[i] = mj
            usados.add(mj)
    for i, h in enumerate(HEADERS):                       # 3) por posición
        if i in mapa:
            continue
        if i < n and i not in usados:
            mapa[i] = i
            usados.add(i)
            avisos.append(f"Columna {LETRAS[i]}: no se encontró el encabezado '{h[:45]}', "
                          f"se usó la columna por posición.")
        else:
            mapa[i] = None
    return mapa, avisos


def leer_excel(archivo):
    contenido = archivo.getvalue() if hasattr(archivo, "getvalue") else archivo
    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    elegida = None
    try:
        for ws in wb.worksheets:
            try:
                ws.reset_dimensions()
            except Exception:
                pass
            filas = [tuple(r) for r in ws.iter_rows(values_only=True)]
            if not filas:
                continue
            h = buscar_fila_encabezado(filas)
            if h is not None:
                elegida = (ws.title, filas, h)
                break
            if elegida is None:
                elegida = (ws.title, filas, None)
    finally:
        wb.close()
    if elegida is None:
        raise ValueError("El archivo no contiene datos.")
    hoja, filas, h = elegida
    if h is None:
        if len(filas[0]) >= 26:
            h = 0
        else:
            raise ValueError(
                "No se encontraron los encabezados esperados (PER_ID, CODIGO_DANE_2025, "
                "NOMBRE_ESTABLECIMIENTO_2025, ...) y el archivo tiene menos de 26 columnas.")
    enc = list(filas[h])
    while enc and (enc[-1] is None or str(enc[-1]).strip() == ""):
        enc.pop()
    n = len(enc)
    mapa, avisos = mapear_columnas(enc)
    faltantes = [HEADERS[i] for i in (0, 2, 3) if mapa.get(i) is None]
    if faltantes:
        raise ValueError("Faltan columnas obligatorias: " + ", ".join(faltantes))
    for i in range(N_INFO):
        if mapa[i] is None:
            avisos.append(f"Columna {LETRAS[i]} ({HEADERS[i]}) no existe en el archivo; se mostrará vacía.")
    for i in range(N_INFO, 26):
        if mapa[i] is None:
            avisos.append(f"Columna {LETRAS[i]} no existe en el archivo; se agregará en la exportación.")

    registros = []
    for pos, fila in enumerate(filas[h + 1:]):
        vals = list(fila[:n]) + [None] * max(0, n - len(fila))
        if not any(v is not None and str(v).strip() != "" for v in vals):
            continue                                       # fila vacía
        orig = []
        for i in range(N_INFO):
            j = mapa[i]
            orig.append(a_texto(vals[j]) if j is not None else "")
        edits = {}
        for i in range(N_INFO, 26):
            j = mapa[i]
            edits[i] = a_texto(vals[j]) if j is not None else ""
        rec = dict(fila=h + 2 + pos, vals=vals, orig=orig, edits=edits, init=dict(edits))
        rec["inst"] = orig[3]
        rec["inst_n"] = norm(orig[3])
        rec["dane"] = orig[2]
        rec["busq"] = norm(" ".join([orig[0], orig[7], orig[8], orig[9], orig[10], orig[11], orig[3]]))
        rec["estado"] = estado_seguimiento(edits)
        registros.append(rec)
    if not registros:
        raise ValueError("El archivo no contiene registros de estudiantes debajo de los encabezados.")
    return dict(hoja=hoja, enc=enc, mapa=mapa, avisos=avisos, registros=registros)


def escribir_xlsx(datos_archivo, indices):
    """Genera el Excel formateado en memoria y devuelve sus bytes."""
    cols = columnas_export(datos_archivo)
    registros = datos_archivo["registros"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Seguimiento"
    ws.append([limpiar(h) for _, _, h in cols])
    for k in indices:
        r = registros[k]
        fila = []
        for tipo, ref, _ in cols:
            if tipo == "canon":
                if ref >= N_INFO:
                    v = r["edits"].get(ref, "")
                else:
                    j = datos_archivo["mapa"].get(ref)
                    v = r["vals"][j] if j is not None and j < len(r["vals"]) else ""
                if ref in COLS_TEXTO:
                    v = a_texto(v)
            else:
                v = r["vals"][ref] if ref < len(r["vals"]) else None
            if isinstance(v, str) and v == "":
                v = None
            fila.append(limpiar(v))
        ws.append(fila)
    fino = Side(style="thin", color="C9D4E3")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    for c in range(1, len(cols) + 1):
        cel = ws.cell(1, c)
        info = c <= N_INFO
        cel.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        cel.fill = PatternFill("solid", fgColor="0B2A5B" if info else ("1E7E4F" if c <= 26 else "5B6B82"))
        cel.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cel.border = borde
        letra = get_column_letter(c)
        if c <= N_INFO:
            mx = max((len(str(ws.cell(rr, c).value or "")) for rr in range(1, min(ws.max_row, 300) + 1)), default=10)
            ws.column_dimensions[letra].width = min(max(mx * 0.95 + 2, 14), 38)
        elif c <= 26:
            ws.column_dimensions[letra].width = 42
        else:
            ws.column_dimensions[letra].width = 22
    ws.row_dimensions[1].height = 92
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(cols)):
        for cel in fila:
            cel.border = borde
            if cel.column in (c + 1 for c in COLS_TEXTO):
                cel.number_format = "@"
            cel.alignment = Alignment(wrap_text=cel.column > N_INFO and cel.column <= 26, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


def columnas_export(datos_archivo):
    cols = []
    usadas = set()
    for i in range(26):
        j = datos_archivo["mapa"].get(i)
        if j is not None:
            usadas.add(j)
            h = datos_archivo["enc"][j] if datos_archivo["enc"][j] is not None else HEADERS[i]
        else:
            h = HEADERS[i]
        cols.append(("canon", i, str(h)))
    for j, h in enumerate(datos_archivo["enc"]):
        if j not in usadas:
            cols.append(("src", j, str(h) if h is not None else ""))
    return cols


def guardar_avance_json(datos_archivo):
    cambios = {}
    for r in datos_archivo["registros"]:
        if r["edits"] != r["init"]:
            cambios[str(r["fila"])] = {"per_id": r["orig"][0], "edits": {str(i): v for i, v in r["edits"].items()}}
    return json.dumps({"version": 2, "archivo": datos_archivo.get("nombre_archivo", "archivo.xlsx"),
                       "guardado": datetime.datetime.now().isoformat(timespec="seconds"),
                       "registros": cambios}, ensure_ascii=False, indent=2).encode("utf-8")


def aplicar_avance(datos_archivo, pack):
    cambios = pack.get("registros", {}) if isinstance(pack, dict) else {}
    por_fila = {str(r["fila"]): r for r in datos_archivo["registros"]}
    aplicados = 0
    for fila, item in cambios.items():
        r = por_fila.get(str(fila))
        if not r or str(item.get("per_id", "")) != str(r["orig"][0]):
            continue
        for i, valor in item.get("edits", {}).items():
            try:
                idx = int(i)
            except (TypeError, ValueError):
                continue
            if idx in range(N_INFO, 26):
                r["edits"][idx] = str(valor or "")
        r["estado"] = estado_seguimiento(r["edits"])
        aplicados += 1
    return aplicados


def sincronizar_widgets():
    datos = st.session_state.get("datos")
    if not datos:
        return
    for k, r in enumerate(datos["registros"]):
        for campo in CAMPOS:
            idx = campo["i"]
            key = f"campo_{k}_{idx}"
            if key in st.session_state:
                r["edits"][idx] = str(st.session_state[key] or "")
        r["estado"] = estado_seguimiento(r["edits"])


def etiqueta_estudiante(r, k):
    nombre = " ".join(x for x in (r["orig"][8], r["orig"][9], r["orig"][10], r["orig"][11]) if x).strip()
    nombre = nombre or "Estudiante sin nombre"
    estado = {0: "Pendiente", 1: "En proceso", 2: "Completado"}[r["estado"]]
    return f"{nombre} · Doc. {r['orig'][7] or 'Sin documento'} · {estado}"


def main():
    st.set_page_config(page_title=APP_TITLE, page_icon="📚", layout="wide")
    st.markdown("""
    <style>
    :root {
      --navy:#07376b; --blue:#0874d9; --blue2:#0d65bc;
      --pale:#edf6ff; --line:#d5e5f5; --ink:#12345b;
      --muted:#647b96; --green:#078344; --gold:#f2a900;
    }
    .stApp { background:#f4f9ff; color:var(--ink); }
    .block-container { padding:0 1rem 1rem !important; max-width:100% !important; }
    header[data-testid="stHeader"] { height:0; background:transparent; }
    [data-testid="stToolbar"] { right:1rem; }
    .hero {
      position:relative; overflow:hidden; display:flex; align-items:center;
      gap:18px; min-height:98px; padding:12px 24px; margin:0 -1rem 12px;
      color:#fff; background:linear-gradient(105deg,#07366a 0%,#07599c 58%,#0879c7 100%);
      border-bottom:4px solid #0b4d88; box-shadow:0 3px 9px #0b3e6c22;
    }
    .hero:after { content:""; position:absolute; right:13%; top:-55px; width:260px;
      height:210px; transform:skewX(-36deg); background:#ffffff0d;
      border-left:2px solid #ffffff14; border-right:2px solid #ffffff10; }
    .hero-icon { font-size:43px; line-height:1; width:72px; text-align:center; z-index:1; }
    .hero-main { z-index:1; flex:1; }
    .hero h1 { font-size:clamp(1.05rem,1.65vw,1.55rem); letter-spacing:.1px;
      margin:0 0 6px; color:white; font-weight:800; line-height:1.2; }
    .hero p { margin:0; color:#e2efff; font-size:.94rem; font-weight:600; }
    .hero-sign { z-index:1; min-width:190px; text-align:center; line-height:1.35; }
    .hero-sign strong { display:block; font-size:.82rem; }
    .hero-sign em { display:block; font-family:cursive; font-size:1.35rem; }
    .hero-sign span { font-size:.75rem; font-weight:700; }
    [data-testid="stSidebar"] { background:linear-gradient(180deg,#eaf5ff,#f4f9ff);
      border-right:1px solid #d3e5f6; }
    [data-testid="stSidebar"] > div { padding-top:.85rem; }
    [data-testid="stSidebar"] h2,[data-testid="stSidebar"] h3 { color:#07376b; }
    h2,h3 { color:#07376b !important; }
    h2 { font-size:1.15rem !important; }
    h3 { font-size:1rem !important; }
    div[data-testid="stMetric"] { background:linear-gradient(145deg,#fff,#f0f7ff);
      border:1px solid #d5e5f5; padding:10px 13px; border-radius:5px;
      box-shadow:0 2px 6px #174a7820; }
    div[data-testid="stMetricLabel"] { color:#17416f; font-weight:700; font-size:.82rem; }
    div[data-testid="stMetricValue"] { color:#075aa5; font-size:1.55rem; font-weight:800; }
    div[data-testid="stTabs"] button { color:#0a396c; font-weight:750; }
    div[data-testid="stTabs"] button[aria-selected="true"] { color:#0874d9; border-bottom-color:#0874d9; }
    div[data-testid="stTextInput"] input,div[data-testid="stTextArea"] textarea,
    div[data-testid="stSelectbox"] [data-baseweb="select"] > div {
      border-radius:4px; border-color:#cbdced; background:#fff; color:#143d69; }
    div[data-testid="stTextInput"] label,div[data-testid="stTextArea"] label,
    div[data-testid="stSelectbox"] label { color:#164577; font-size:.79rem; font-weight:650; }
    div.stButton > button,div.stDownloadButton > button { border-radius:4px;
      min-height:2.45rem; font-weight:750; border:1px solid #0a70c8; }
    div.stButton > button[kind="primary"],div.stDownloadButton > button {
      background:linear-gradient(100deg,#0873d1,#075db3); color:#fff; }
    div.stButton > button:hover,div.stDownloadButton > button:hover {
      border-color:#034e91; color:#fff; }
    div[data-testid="stRadio"] label { background:#fff; border:1px solid #dce8f5;
      border-radius:3px; padding:5px 7px; margin-bottom:3px; }
    div[data-testid="stRadio"] label p { font-size:.78rem; line-height:1.2; }
    div[data-testid="stProgress"] > div { background:#dceafa; }
    div[data-testid="stProgress"] div[role="progressbar"] > div { background:#078344; }
    hr { border-color:#d5e5f5; }
    .app-footer { display:flex; justify-content:space-between; align-items:center; gap:12px;
      margin:12px -1rem 0; padding:12px 22px; background:linear-gradient(100deg,#07366a,#07599c);
      color:#fff; font-size:.82rem; border-top:2px solid #0b4d88; }
    .app-footer em { font-family:cursive; font-size:.98rem; }
    @media(max-width:900px) { .hero { padding:12px; gap:10px; }
      .hero-icon,.hero-sign { display:none; } .app-footer { display:block; }
      .app-footer em { display:block; margin-top:5px; } }
    </style>
    """, unsafe_allow_html=True)
    st.markdown(f'<div class="hero"><div class="hero-icon">📚</div><div class="hero-main"><h1>{APP_TITLE}</h1><p>{SUBTITULO}</p></div><div class="hero-sign"><strong>Creado por:</strong><em>Jaime Hernandez</em><span>- Montería -</span></div></div>', unsafe_allow_html=True)

    with st.sidebar:
        st.header("Archivo y filtros")
        excel = st.file_uploader("Cargar archivo Excel (.xlsx)", type=["xlsx"], key="uploader_excel")
        avance = st.file_uploader("Cargar avance anterior (.json), opcional", type=["json"], key="uploader_json")
        if excel is not None:
            token = (excel.name, len(excel.getvalue()))
            if st.session_state.get("archivo_token") != token:
                try:
                    datos = leer_excel(io.BytesIO(excel.getvalue()))
                    datos["nombre_archivo"] = excel.name
                    st.session_state["datos"] = datos
                    st.session_state["archivo_token"] = token
                    st.session_state.pop("seleccion_estudiante", None)
                    st.session_state["avance_aplicado_token"] = None
                    st.success(f"Archivo cargado: {len(datos['registros']):,} registros")
                except Exception as exc:
                    st.error(f"No se pudo leer el Excel: {exc}")
                    st.stop()
        datos = st.session_state.get("datos")
        if datos and avance is not None:
            token_json = (avance.name, len(avance.getvalue()))
            if st.session_state.get("avance_aplicado_token") != token_json:
                try:
                    pack = json.loads(avance.getvalue().decode("utf-8-sig"))
                    n = aplicar_avance(datos, pack)
                    st.session_state["avance_aplicado_token"] = token_json
                    st.success(f"Avance aplicado a {n:,} registros compatibles")
                except Exception as exc:
                    st.error(f"No se pudo leer el avance JSON: {exc}")
        if not datos:
            st.info("Cargue un archivo Excel para iniciar el seguimiento.")
            return

        sincronizar_widgets()
        registros = datos["registros"]
        instituciones = sorted({r["inst"] for r in registros if r["inst"]}, key=str.casefold)
        filtro_inst = st.selectbox("Institución Educativa", [TODAS] + instituciones, key="filtro_inst")
        danes = sorted({r["dane"] for r in registros if r["dane"]}, key=str.casefold)
        filtro_dane = st.selectbox("Código DANE", ["Todos"] + danes, key="filtro_dane")
        busqueda = st.text_input("Buscar estudiante", placeholder="Nombre, documento o PER_ID", key="busqueda").strip()

    sincronizar_widgets()
    registros = datos["registros"]
    alcance = []
    for k, r in enumerate(registros):
        if filtro_inst != TODAS and r["inst"] != filtro_inst:
            continue
        if filtro_dane != "Todos" and r["dane"] != filtro_dane:
            continue
        if busqueda and norm(busqueda) not in r["busq"]:
            continue
        alcance.append(k)

    total = len(alcance)
    completos = sum(registros[k]["estado"] == 2 for k in alcance)
    proceso = sum(registros[k]["estado"] == 1 for k in alcance)
    pendientes = sum(registros[k]["estado"] == 0 for k in alcance)
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Total registros", f"{total:,}")
    m2.metric("Completados", f"{completos:,}")
    m3.metric("En proceso", f"{proceso:,}")
    m4.metric("Pendientes", f"{pendientes:,}")
    if total:
        st.progress(completos / total, text=f"Avance completo: {completos:,} de {total:,} ({completos / total:.1%})")
    st.divider()
    izquierda, derecha = st.columns([0.95, 1.6], gap="large")
    with izquierda:
        st.subheader("Estudiantes")
        st.caption(f"Mostrando {total:,} registros según los filtros")
        if not alcance:
            st.warning("No hay estudiantes que coincidan con los filtros.")
            st.session_state.pop("seleccion_estudiante", None)
        else:
            if st.session_state.get("seleccion_estudiante") not in alcance:
                st.session_state["seleccion_estudiante"] = alcance[0]
            elegido = st.radio("Seleccione un estudiante", options=alcance,
                               format_func=lambda k: etiqueta_estudiante(registros[k], k),
                               key="seleccion_estudiante", label_visibility="collapsed")
            st.caption("Estados: Completado · En proceso · Pendiente")
    with derecha:
        st.subheader("Ficha de seguimiento")
        if alcance:
            r = registros[elegido]
            st.markdown(f"**Institución:** {r['inst'] or 'Sin información'}  \n**Código DANE:** {r['dane'] or 'Sin información'}")
            tab1, tab2 = st.tabs(["Datos básicos SIMAT · A:N", "Seguimiento · O:Z"])
            with tab1:
                campos_info = [("PER_ID", 0), ("Año", 1), ("Código DANE", 2), ("Institución educativa", 3),
                               ("Código DANE sede", 4), ("Sede", 5), ("Tipo documento", 6), ("Número documento", 7),
                               ("Primer apellido", 8), ("Segundo apellido", 9), ("Primer nombre", 10), ("Segundo nombre", 11),
                               ("Sector", 12), ("Estado definitivo", 13)]
                for start in range(0, len(campos_info), 2):
                    c1, c2 = st.columns(2)
                    for col, (label, idx) in zip((c1, c2), campos_info[start:start+2]):
                        with col:
                            st.text_input(label, value=r["orig"][idx], key=f"info_{elegido}_{idx}", disabled=True)
            with tab2:
                st.caption("Complete los campos de seguimiento. Los datos básicos SIMAT permanecen bloqueados.")
                for campo in CAMPOS:
                    idx = campo["i"]
                    key = f"campo_{elegido}_{idx}"
                    valor = str(r["edits"].get(idx, "") or "")
                    if campo["tipo"] == "combo":
                        opciones = list(campo.get("op", []))
                        if not campo.get("cerrado", False) and valor and valor not in opciones:
                            opciones = [valor] + opciones
                        if not campo.get("cerrado", False):
                            opciones = [""] + opciones
                        else:
                            opciones = [""] + opciones
                        actual = valor if valor in opciones else ""
                        st.selectbox(f"{LETRAS[idx]} · {campo['corto']}", opciones,
                                     index=opciones.index(actual), key=key)
                    else:
                        st.text_area(f"{LETRAS[idx]} · {campo['corto']}", value=valor,
                                     height=100 if campo.get("alto", 4) >= 4 else 75, key=key)
                sincronizar_widgets()
                estado = r["estado"]
                st.info({0: "Estado del registro: Pendiente", 1: "Estado del registro: En proceso", 2: "Estado del registro: Completado"}[estado])

    st.divider()
    st.subheader("Descargas oficiales")
    sincronizar_widgets()
    registros = datos["registros"]
    indices_export = list(alcance)
    json_bytes = guardar_avance_json(datos)
    nombre_base = "consolidado" if filtro_inst == TODAS else slug(filtro_inst)
    nombre_excel = nombre_base + SUFIJO + ".xlsx"
    cjson, cxlsx = st.columns(2)
    with cjson:
        st.download_button("Descargar avance (.json)", data=json_bytes,
                           file_name=(nombre_base + "_avance.json"), mime="application/json",
                           use_container_width=True, help="Guarda únicamente los cambios para poder retomarlos después.")
    with cxlsx:
        if indices_export:
            excel_bytes = escribir_xlsx(datos, indices_export)
            st.download_button("Exportar Excel formateado (.xlsx)", data=excel_bytes,
                               file_name=nombre_excel, mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               use_container_width=True, help="Exporta los registros que coinciden con los filtros actuales.")
        else:
            st.button("Exportar Excel formateado (.xlsx)", disabled=True, use_container_width=True)
    st.markdown(f'<div class="app-footer"><span>{PIE}</span><em>{LEMA}</em></div>', unsafe_allow_html=True)
    st.caption("Verifique la información antes de compartir o remitir los archivos.")


if __name__ == "__main__":
    main()
